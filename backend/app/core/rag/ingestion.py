import uuid
from pathlib import Path
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct
from langchain_openai import OpenAIEmbeddings

from app.config import settings

COLLECTION = settings.qdrant_collection


def _get_embeddings():
    if settings.default_llm_provider in ("doubao",):
        return OpenAIEmbeddings(
            model=settings.embedding_model,
            api_key=settings.doubao_api_key,
            base_url=settings.doubao_base_url,
            dimensions=settings.embedding_dim,
        )
    return OpenAIEmbeddings(
        model=settings.embedding_model,
        api_key=settings.openai_api_key,
        dimensions=settings.embedding_dim,
    )


def _get_qdrant() -> QdrantClient:
    if settings.use_local_qdrant:
        return QdrantClient(path=settings.qdrant_local_path)
    return QdrantClient(url=settings.qdrant_url)


def ensure_collection():
    client = _get_qdrant()
    collections = [c.name for c in client.get_collections().collections]
    if COLLECTION not in collections:
        client.create_collection(
            collection_name=COLLECTION,
            vectors_config=VectorParams(
                size=settings.embedding_dim,
                distance=Distance.COSINE,
            ),
        )


async def ingest_document(
    content: str,
    filename: str,
    doc_type: str = "text",
    metadata: dict | None = None,
) -> int:
    from app.core.rag.chunking import chunk_document

    ensure_collection()
    embeddings = _get_embeddings()
    client = _get_qdrant()

    chunks = chunk_document(content, doc_type)
    if not chunks:
        return 0

    doc_id = str(uuid.uuid4())
    points = []
    for i, chunk in enumerate(chunks):
        vector = embeddings.embed_query(chunk)
        points.append(PointStruct(
            id=str(uuid.uuid4()),
            vector=vector,
            payload={
                "doc_id": doc_id,
                "filename": filename,
                "chunk_index": i,
                "content": chunk,
                "doc_type": doc_type,
                **(metadata or {}),
            },
        ))

    client.upsert(collection_name=COLLECTION, points=points)
    return len(points)


async def ingest_file(file_path: Path) -> int:
    suffix = file_path.suffix.lower()
    type_map = {
        ".txt": "text",
        ".md": "markdown",
        ".pdf": "pdf",
        ".docx": "docx",
        ".xlsx": "xlsx",
        ".pptx": "pptx",
    }
    doc_type = type_map.get(suffix, "text")

    if suffix == ".pdf":
        import fitz
        doc = fitz.open(file_path)
        content = "\n".join(page.get_text() for page in doc)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(file_path)
        content = "\n".join(p.text for p in doc.paragraphs)
    elif suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(file_path)
        parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join(str(c) if c else "" for c in row))
            parts.append(f"## {sheet}\n" + "\n".join(rows))
        content = "\n\n".join(parts)
    else:
        content = file_path.read_text(encoding="utf-8")

    return await ingest_document(content, file_path.name, doc_type)
