import re
from fastapi import APIRouter, UploadFile, File, HTTPException
from pydantic import BaseModel

from app.config import settings

router = APIRouter(prefix="/api/knowledge", tags=["knowledge"])

_doc_store: dict[str, list[str]] = {}
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100


def _split_text(text: str) -> list[str]:
    paragraphs = re.split(r'\n\s*\n', text)
    chunks = []
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(para) <= CHUNK_SIZE:
            chunks.append(para)
        else:
            sentences = re.split(r'(?<=[。！？.!?])\s*', para)
            current = ""
            for sent in sentences:
                if len(current) + len(sent) <= CHUNK_SIZE:
                    current += sent
                else:
                    if current:
                        chunks.append(current.strip())
                    current = sent
                    while len(current) > CHUNK_SIZE:
                        chunks.append(current[:CHUNK_SIZE])
                        current = current[CHUNK_SIZE - CHUNK_OVERLAP:]
            if current.strip():
                chunks.append(current.strip())
    return chunks


def _retrieve_chunks(query: str, chunks: list[str], top_k: int = 5) -> list[str]:
    query_words = set(re.findall(r'[一-鿿]+|[a-zA-Z]+', query.lower()))
    if not query_words:
        return chunks[:top_k]

    scored = []
    for i, chunk in enumerate(chunks):
        chunk_lower = chunk.lower()
        score = sum(chunk_lower.count(w) for w in query_words)
        if score > 0:
            scored.append((score, i))
    scored.sort(key=lambda x: -x[0])
    top_indices = [i for _, i in scored[:top_k]]
    if not top_indices:
        return chunks[:top_k]
    top_indices.sort()
    return [chunks[i] for i in top_indices]


class IngestTextRequest(BaseModel):
    content: str
    filename: str
    doc_type: str = "text"


class IngestResponse(BaseModel):
    status: str
    filename: str
    content_preview: str
    size: int
    chunks: int = 0


@router.post("/ingest/text")
async def ingest_text(req: IngestTextRequest) -> IngestResponse:
    chunks = _split_text(req.content)
    _doc_store[req.filename] = chunks
    preview = req.content[:300] + ("..." if len(req.content) > 300 else "")
    return IngestResponse(status="ok", filename=req.filename, content_preview=preview, size=len(req.content), chunks=len(chunks))


@router.post("/ingest/file")
async def ingest_upload(file: UploadFile = File(...)) -> IngestResponse:
    if file.size and file.size > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(400, f"文件大小超过 {settings.max_upload_size_mb}MB 限制")

    content = await file.read()
    suffix = file.filename.split(".")[-1].lower() if "." in file.filename else ""

    try:
        if suffix in ("txt", "md"):
            text = content.decode("utf-8")
        elif suffix == "pdf":
            import fitz
            doc = fitz.open(stream=content, filetype="pdf")
            text = "\n".join(page.get_text() for page in doc)
        elif suffix == "docx":
            from docx import Document
            import io as _io
            doc = Document(_io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
        elif suffix == "xlsx":
            import openpyxl
            import io as _io
            wb = openpyxl.load_workbook(_io.BytesIO(content))
            parts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c else "" for c in row))
                parts.append(f"## {sheet}\n" + "\n".join(rows))
            text = "\n\n".join(parts)
        elif suffix == "pptx":
            from pptx import Presentation
            import io as _io
            prs = Presentation(_io.BytesIO(content))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                lines = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                lines.append(t)
                if lines:
                    slides.append(f"## 幻灯片 {i}\n" + "\n".join(lines))
            text = "\n\n".join(slides)
        else:
            text = content.decode("utf-8")
    except Exception as e:
        raise HTTPException(400, f"文件解析失败: {str(e)}")

    chunks = _split_text(text)
    _doc_store[file.filename] = chunks
    preview = text[:300] + ("..." if len(text) > 300 else "")
    return IngestResponse(status="ok", filename=file.filename, content_preview=preview, size=len(text), chunks=len(chunks))


def get_doc_content(filename: str | None = None, query: str | None = None) -> dict:
    if filename:
        chunks = _doc_store.get(filename, [])
        if query and chunks:
            relevant = _retrieve_chunks(query, chunks, top_k=5)
            return {"filename": filename, "content": "\n\n---\n\n".join(relevant), "chunks_used": len(relevant), "total_chunks": len(chunks)}
        return {"filename": filename, "content": "\n\n---\n\n".join(chunks), "total_chunks": len(chunks)}
    return {"files": list(_doc_store.keys()), "documents": {k: len(v) for k, v in _doc_store.items()}}
